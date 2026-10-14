"""대기압 vs HSOH 두께 분석.

데이터 3개를 연결해서 분석한다.
    쿼리1 대기압      : 시간, 대기압
    쿼리2 계측        : LOT, WAFER, ITEM, 두께 (+ 계측시간)
    쿼리3 공정 진행   : LOT, 설비, TKIN_TIME (+ WAFER)
    쿼리2 + 쿼리3 은 LOT(있으면 WAFER 까지)으로 연결 -> 웨이퍼별 설비, TKIN_TIME
    그 TKIN_TIME 에 가장 가까운 쿼리1 대기압을 붙인다.

사용법 (VS Code 에서 F5 또는 python analyze.py)
    1) 아래 "설정" 에서 기간(DATE_FROM/DATE_TO), SQL 3개, COLUMNS 를 채운다.
    2) 실행하면 bigdataquery.getData() 로 조회 -> 연결 -> 분석 -> output/ 에 리포트 생성.
       조회 원본은 output/raw_pressure.csv, raw_meas.csv, raw_proc.csv 로도 저장된다.

    DATA_SOURCE = "csv" 로 바꾸면 조회 대신 input/pressure/, input/meas/, input/proc/ 의 CSV 를 읽는다.
    명령줄로 덮어쓰기:  python analyze.py --from 2026-09-01 --to 2026-09-30
                        python analyze.py --source csv

조회 결과/CSV 에 컬럼이 많아도 COLUMNS 에 지정한 컬럼만 추출해서 사용한다.
"""
import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd

# ============================================================
# 설정 (실제 데이터에 맞게 수정)
# ============================================================
DATA_SOURCE = "query"            # "query": bigdataquery 로 조회 / "csv": input 폴더 CSV 사용
DATE_FROM = "2026-09-01"         # 분석 시작일 (계측 기준, 포함)
DATE_TO = "2026-09-30"           # 분석 종료일 (계측 기준, 포함)
PROC_LOOKBACK_DAYS = 3           # 쿼리3(공정) 조회를 계측 시작일보다 며칠 앞부터 할지

# 조회 SQL. 아래 자리표시자는 실행할 때 자동으로 바뀐다.
#   {start}, {end}         -> '2026-09-01 00:00:00' 형식 (end 는 종료일 다음날 0시, 미포함)
#   {start_ymd}, {end_ymd} -> '20260901' 형식 (날짜 파티션 컬럼용, 둘 다 포함)
#   {lots}                 -> 쿼리2 에서 나온 LOT 목록 'LOT1','LOT2',... (쿼리3 에서만 사용 가능)
# 쿼리1(대기압) 의 기간은 쿼리3 에서 나온 TKIN_TIME 범위에 맞춰 자동으로 정해진다.
# 쿼리3 의 {start} 는 계측 시작일 - PROC_LOOKBACK_DAYS 이다.

# 쿼리1: 대기압
SQL_PRESSURE = """
select *
from 대기압쿼리명
where 1=1
  and act_time >= '{start}'
  and act_time <  '{end}'
"""

# 쿼리2: LOT - 계측 데이터
SQL_MEAS = """
select *
from 계측쿼리명
where 1=1
  and ITEM_ID = 'AVG'
  and MEAS_TIME >= '{start}'
  and MEAS_TIME <  '{end}'
"""

# 쿼리3: LOT - 공정 진행 정보 (HSOH 공정만 나오도록 STEP 조건을 꼭 넣을 것)
#   {start} 는 계측 시작일 - PROC_LOOKBACK_DAYS (계측 전에 진행된 공정까지 포함하려고)
#   기간 대신 LOT 으로 조회하려면:  and LOT_ID in ({lots})
SQL_PROC = """
select *
from 공정진행쿼리명
where 1=1
  and TKIN_TIME >= '{start}'
  and TKIN_TIME <  '{end}'
  -- and STEP_ID = 'HSOH 공정 STEP'
"""

# 각 쿼리 결과에서 사용할 컬럼명. None 은 "그 컬럼 없음" 이라는 뜻.
COLUMNS = {
    "pressure": {                # 쿼리1
        "time": "act_time",      # 대기압 측정 시간
        "value": "value",        # 대기압 값 (hPa)
    },
    "meas": {                    # 쿼리2
        "lot": "LOT_ID",
        "wafer": "WAFER_ID",
        "item": "ITEM_ID",
        "value": "FAB_VALUE",    # 두께 (Å)
        "meas_time": "MEAS_TIME",  # 계측 시간 (없으면 None)
    },
    "proc": {                    # 쿼리3
        "lot": "LOT_ID",
        "eqp": "EQP_ID",
        "time": "TKIN_TIME",     # 공정 진행 시간 -> 대기압 매칭 기준
        "wafer": None,           # 공정 정보가 WAFER 단위면 컬럼명 입력 (LOT 단위면 None)
    },
}
ITEM_FILTER = "AVG"              # ITEM_ID 가 이 값인 행만 사용 (None 이면 필터 안 함)
MATCH_TOLERANCE_MIN = 5          # TKIN_TIME 과 대기압 시간 차이 허용 범위(분)
WINDOW_MIN = 30                  # 직전 평균/변화량 계산 구간(분)
PRESSURE_BIN_HPA = 1.0           # 대기압 구간 폭(hPa)
LAG_MAX_MIN = 240                # 시차 분석 최대 범위(분)
LAG_STEP_MIN = 15                # 시차 분석 간격(분)
MIN_N_FOR_STATS = 3              # 통계 계산 최소 웨이퍼 수
EXCEL_MAX_ROWS = 20000           # 매칭데이터가 이보다 많으면 엑셀 대신 CSV 로 저장 (엑셀 저장이 느려짐)
MAKE_PNG = True                  # False 면 PNG 그래프 생략 (HTML 리포트로 충분할 때 시간 절약)

ENCODINGS = ("utf-8-sig", "cp949", "euc-kr", "utf-8")
ROOT = Path(__file__).resolve().parent if "__file__" in globals() else Path.cwd()
SOURCES = {"pressure": "쿼리1(대기압)", "meas": "쿼리2(계측)", "proc": "쿼리3(공정)"}


# ============================================================
# 읽기
# ============================================================
def _norm(name):
    """컬럼명 비교용: 대소문자, 공백, '_', '-', '.' 차이는 무시 (LOT ID = lot_id = LOT-ID)."""
    return re.sub(r"[\s_\-.\ufeff]+", "", str(name)).upper()


def _wanted(kind):
    return {k: v for k, v in COLUMNS[kind].items() if v}


def pick_columns(df, kind, src, all_cols=None):
    """df 에서 COLUMNS[kind] 에 지정한 컬럼만 골라 내부이름으로 바꿔 반환."""
    wanted = _wanted(kind)
    targets = {_norm(v): k for k, v in wanted.items()}
    all_cols = df.columns if all_cols is None else all_cols
    df = df[[c for c in df.columns if _norm(c) in targets]]
    df = df.rename(columns=lambda c: targets[_norm(c)])
    df = df.loc[:, ~df.columns.duplicated()]
    missing = [wanted[k] for k in wanted if k not in df.columns]
    if missing:
        raise SystemExit(f"[오류] {src} 에 필요한 컬럼이 없습니다: {missing}\n"
                         f"       실제 컬럼: {list(all_cols)}\n"
                         f"       analyze.py 상단 COLUMNS 설정(또는 SQL 의 AS 이름)을 확인하세요.")
    return df[list(wanted)].astype("string").apply(lambda x: x.str.strip())


def read_csv_columns(path, kind):
    """CSV 에서 COLUMNS[kind] 컬럼만 읽는다. 인코딩은 자동 판별."""
    targets = {_norm(v) for v in _wanted(kind).values()}
    last_err = None
    for enc in ENCODINGS:
        try:
            header = pd.read_csv(path, encoding=enc, nrows=0).columns
            df = pd.read_csv(path, encoding=enc, dtype=str,
                             usecols=lambda c: _norm(c) in targets)
            break
        except UnicodeDecodeError as e:
            last_err = e
    else:
        raise SystemExit(f"[오류] {path} 인코딩을 읽을 수 없습니다: {last_err}")
    return pick_columns(df, kind, path, header)


def collect_csvs(path):
    path = Path(path)
    if path.is_dir():
        files = sorted(p for p in path.iterdir() if p.suffix.lower() == ".csv")
    elif path.exists():
        files = [path]
    else:
        files = []
    if not files:
        raise SystemExit(f"[오류] CSV 파일이 없습니다: {path}")
    return files


def read_all(path, kind):
    files = collect_csvs(path)
    frames = [read_csv_columns(f, kind) for f in files]
    print(f"  {SOURCES[kind]} {path}: {len(files)}개 파일, {sum(map(len, frames))}행")
    return pd.concat(frames, ignore_index=True)


_KR_TIME = re.compile(
    r"^\s*(\d{4})[-./](\d{1,2})[-./](\d{1,2})\s*(오전|오후|AM|PM|am|pm)\s*"
    r"(\d{1,2}):(\d{2})(?::(\d{2}))?\s*$")


def parse_time(s):
    """'2026-09-14 오후 4:14:59' 형식(및 일반 24시간 형식) 을 datetime 으로 변환."""
    s = s.astype("string")
    parts = s.str.extract(_KR_TIME)
    ok = parts[0].notna()
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    if ok.any():
        p = parts[ok]
        hour = p[4].astype(int) % 12
        hour = hour + np.where(p[3].str.upper().isin(["오후", "PM"]), 12, 0)
        out[ok] = pd.to_datetime(pd.DataFrame({
            "year": p[0].astype(int), "month": p[1].astype(int), "day": p[2].astype(int),
            "hour": hour, "minute": p[5].astype(int),
            "second": p[6].fillna("0").astype(int)}), errors="coerce")
    rest = ~ok & s.notna()
    if rest.any():
        out[rest] = pd.to_datetime(s[rest], errors="coerce")
    return out


def load_pressure(raw):
    df = raw.copy()
    df["time"] = parse_time(df["time"])
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    bad = df["time"].isna() | df["value"].isna()
    if bad.any():
        print(f"  대기압: 시간/값을 읽지 못한 {bad.sum()}행 제외")
    df = df[~bad].drop_duplicates("time").sort_values("time").reset_index(drop=True)
    return df.rename(columns={"value": "pressure"})


def load_meas(raw):
    df = raw.copy()
    if ITEM_FILTER is not None:
        before = len(df)
        df = df[df["item"].str.upper().eq(ITEM_FILTER.upper()).fillna(False)]
        print(f"  계측: ITEM_ID={ITEM_FILTER} 필터 {before} -> {len(df)}행")
    df["thickness"] = pd.to_numeric(df["value"], errors="coerce")
    if "meas_time" in df:
        df["meas_time"] = parse_time(df["meas_time"])
    bad = df["thickness"].isna() | df["lot"].isna() | df["wafer"].isna()
    if bad.any():
        print(f"  계측: LOT/WAFER/두께를 읽지 못한 {bad.sum()}행 제외")
    df = df[~bad].drop(columns=["value", "item"]).drop_duplicates()
    return df.reset_index(drop=True)


def load_proc(raw):
    df = raw.copy()
    df["time"] = parse_time(df["time"])
    bad = df["time"].isna() | df["lot"].isna()
    if bad.any():
        print(f"  공정: LOT/TKIN_TIME 을 읽지 못한 {bad.sum()}행 제외")
    return df[~bad].drop_duplicates().reset_index(drop=True)


# ============================================================
# bigdataquery 조회
# ============================================================
def render_sql(sql, start, end, lots=None):
    """SQL 의 자리표시자를 치환. end 는 미포함."""
    rep = {"{start}": start.strftime("%Y-%m-%d %H:%M:%S"),
           "{end}": end.strftime("%Y-%m-%d %H:%M:%S"),
           "{start_ymd}": start.strftime("%Y%m%d"),
           "{end_ymd}": (end - pd.Timedelta(seconds=1)).strftime("%Y%m%d")}
    if lots is not None:
        rep["{lots}"] = ",".join("'" + str(x).replace("'", "''") + "'" for x in lots) or "''"
    for k, v in rep.items():
        sql = sql.replace(k, v)
    return sql


def run_query(sql, name):
    try:
        from bigdataquery import getData
    except ImportError:
        raise SystemExit("[오류] bigdataquery 를 불러올 수 없습니다. "
                         "평소 getData 를 쓰던 파이썬 환경(인터프리터)에서 실행하세요.")
    try:
        df = getData(sql)
    except Exception as e:
        raise SystemExit(f"[오류] {name} 조회 실패: {e}\n--- SQL ---\n{sql}")
    if df is None:
        raise SystemExit(f"[오류] {name} 조회 결과가 없습니다.\n--- SQL ---\n{sql}")
    df = df if isinstance(df, pd.DataFrame) else pd.DataFrame(df)
    print(f"  {name}: {len(df)}행 조회")
    return df


def _fetch(kind, sql, start, end, out_dir, lots=None):
    df = run_query(render_sql(sql, start, end, lots), SOURCES[kind])
    df.to_csv(out_dir / f"raw_{kind}.csv", index=False, encoding="utf-8-sig")
    return pick_columns(df, kind, f"{SOURCES[kind]} 조회 결과")


def fetch_from_query(start, end, out_dir):
    """쿼리2 -> 쿼리3 -> 쿼리1 순서로 조회 (앞 결과로 다음 조회 범위를 정한다)."""
    meas = _fetch("meas", SQL_MEAS, start, end, out_dir)
    if meas.empty:
        raise SystemExit("[오류] 쿼리2(계측) 조회 결과가 0행입니다. 기간/조건을 확인하세요.")
    lots = sorted(meas["lot"].dropna().str.strip().unique())
    proc = _fetch("proc", SQL_PROC, start - pd.Timedelta(days=PROC_LOOKBACK_DAYS), end,
                  out_dir, lots)
    if proc.empty:
        raise SystemExit("[오류] 쿼리3(공정) 조회 결과가 0행입니다. 기간/조건을 확인하세요.")

    # 대기압은 (계측 LOT 의) TKIN_TIME 범위 + 시차 분석/직전 평균에 필요한 앞부분까지 조회
    rel = proc[proc["lot"].isin(lots)]
    t = parse_time((rel if len(rel) else proc)["time"]).dropna()
    if t.empty:
        raise SystemExit("[오류] 쿼리3(공정) 의 TKIN_TIME 을 읽지 못했습니다. 시간 형식을 확인하세요.")
    margin = pd.Timedelta(minutes=max(LAG_MAX_MIN, WINDOW_MIN) + MATCH_TOLERANCE_MIN)
    p_start = (t.min() - margin).floor("min")
    p_end = (t.max() + pd.Timedelta(minutes=MATCH_TOLERANCE_MIN + 1)).ceil("min")
    pressure = _fetch("pressure", SQL_PRESSURE, p_start, p_end, out_dir)
    return pressure, meas, proc


# ============================================================
# 연결 (쿼리2 + 쿼리3)
# ============================================================
def join_meas_proc(meas, proc):
    """계측에 공정 정보(설비, TKIN_TIME)를 붙인다.

    LOT(공정 정보에 WAFER 가 있으면 LOT+WAFER)으로 연결한다.
    한 웨이퍼에 공정 이력이 여러 건이면 계측 시간 이전의 가장 늦은 TKIN_TIME 을 쓴다.
    (계측 시간이 없으면 가장 늦은 TKIN_TIME)
    """
    keys = ["lot", "wafer"] if "wafer" in proc else ["lot"]
    m = meas.reset_index(names="_mid")
    j = m.merge(proc, on=keys, how="inner")
    n_multi = (j.groupby("_mid").size() > 1).sum()
    if "meas_time" in j:
        j = j[~(j["meas_time"].notna() & (j["time"] > j["meas_time"]))]
    j = j.sort_values("time").drop_duplicates("_mid", keep="last")
    if n_multi:
        print(f"  공정 이력이 2건 이상인 웨이퍼 {n_multi}개 -> 계측 이전 가장 늦은 TKIN_TIME 사용")
    joined = j.sort_values("_mid").drop(columns="_mid").reset_index(drop=True)
    missing = m[~m["_mid"].isin(j["_mid"])].drop(columns="_mid").reset_index(drop=True)
    return joined, missing


# ============================================================
# 대기압 매칭
# ============================================================
def _asof(left_time, pressure, col, tol_min, direction="nearest"):
    left = pd.DataFrame({"t": left_time.values, "_i": np.arange(len(left_time))}).sort_values("t")
    m = pd.merge_asof(left, pressure[["time", col]].rename(columns={"time": "t"}),
                      on="t", direction=direction,
                      tolerance=pd.Timedelta(minutes=tol_min))
    return m.sort_values("_i")[col].to_numpy()


def match(soh, pressure):
    p = pressure.set_index("time")
    roll = p["pressure"].rolling(f"{WINDOW_MIN}min").mean().rename("roll").reset_index()

    df = soh.copy()
    df["pressure"] = _asof(df["time"], pressure, "pressure", MATCH_TOLERANCE_MIN)
    tmp = pressure.rename(columns={"time": "p_time"}).assign(time=pressure["time"])
    df["pressure_time"] = _asof(df["time"], tmp, "p_time", MATCH_TOLERANCE_MIN)
    df[f"pressure_avg_{WINDOW_MIN}m"] = _asof(df["time"], roll, "roll",
                                              MATCH_TOLERANCE_MIN, "backward")
    before = _asof(df["time"] - pd.Timedelta(minutes=WINDOW_MIN), pressure, "pressure",
                   MATCH_TOLERANCE_MIN)
    df[f"pressure_change_{WINDOW_MIN}m"] = df["pressure"] - before
    return df


# ============================================================
# 분석
# ============================================================
def regress(x, y):
    from scipy import stats
    m = x.notna() & y.notna()
    x, y = x[m].astype(float), y[m].astype(float)
    n = len(x)
    row = {"N": n, "기울기(Å/hPa)": np.nan, "절편": np.nan, "R²": np.nan,
           "Pearson r": np.nan, "p-value": np.nan, "Spearman ρ": np.nan,
           "두께 평균(Å)": y.mean() if n else np.nan, "두께 σ(Å)": y.std() if n > 1 else np.nan,
           "대기압 min": x.min() if n else np.nan, "대기압 max": x.max() if n else np.nan}
    if n >= MIN_N_FOR_STATS and x.nunique() > 1:
        r = stats.linregress(x, y)
        row.update({"기울기(Å/hPa)": r.slope, "절편": r.intercept, "R²": r.rvalue ** 2,
                    "Pearson r": r.rvalue, "p-value": r.pvalue,
                    "Spearman ρ": stats.spearmanr(x, y).statistic})
    return row


def summary_table(df):
    rows = [{"구분": "전체 (두께 원값)", **regress(df["pressure"], df["thickness"])},
            {"구분": "전체 (설비 평균 제거)", **regress(df["pressure"], df["thickness_dev"])}]
    for eqp, g in df.groupby("eqp"):
        rows.append({"구분": f"설비 {eqp}", **regress(g["pressure"], g["thickness"])})
    return pd.DataFrame(rows)


def bin_table(df):
    lo = np.floor(df["pressure"].min() / PRESSURE_BIN_HPA) * PRESSURE_BIN_HPA
    hi = np.ceil(df["pressure"].max() / PRESSURE_BIN_HPA) * PRESSURE_BIN_HPA + PRESSURE_BIN_HPA
    edges = np.arange(lo, hi + 1e-9, PRESSURE_BIN_HPA)
    d = df.assign(대기압구간=pd.cut(df["pressure"], edges, right=False))
    d["대기압구간"] = d["대기압구간"].astype(str)
    agg = {"thickness": ["count", "mean", "std", "min", "max"], "thickness_dev": ["mean"]}
    all_ = d.groupby("대기압구간", sort=False).agg(agg)
    all_.insert(0, "설비", "전체")
    per = d.groupby(["eqp", "대기압구간"], sort=False).agg(agg).reset_index(level=0)
    per = per.rename(columns={"eqp": "설비"})
    out = pd.concat([all_, per])
    out.columns = ["설비", "웨이퍼수", "두께 평균", "두께 σ", "두께 min", "두께 max", "설비편차 평균"]
    out = out.reset_index().rename(columns={"index": "대기압구간"})
    out["_k"] = out["대기압구간"].str.extract(r"([\d.]+)").astype(float)
    return out.sort_values(["설비", "_k"]).drop(columns="_k").reset_index(drop=True)


def lag_table(soh, pressure):
    rows = []
    for lag in range(0, LAG_MAX_MIN + 1, LAG_STEP_MIN):
        p = _asof(soh["time"] - pd.Timedelta(minutes=lag), pressure, "pressure",
                  MATCH_TOLERANCE_MIN)
        dev = soh["thickness"] - soh.groupby("eqp")["thickness"].transform("mean")
        r = regress(pd.Series(p, index=soh.index), dev)
        rows.append({"시차(분, TKIN 이전)": lag, "N": r["N"], "Pearson r": r["Pearson r"],
                     "R²": r["R²"], "기울기(Å/hPa)": r["기울기(Å/hPa)"]})
    return pd.DataFrame(rows)


# ============================================================
# 그래프
# ============================================================
def make_charts(df, pressure, summary, lag, out_dir):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams.update({"figure.dpi": 110, "axes.grid": True, "grid.alpha": 0.3,
                         "axes.spines.top": False, "axes.spines.right": False})
    colors = plt.get_cmap("tab10")
    eqps = sorted(df["eqp"].unique())
    cmap = {e: colors(i % 10) for i, e in enumerate(eqps)}
    files = []

    def save(fig, name):
        p = out_dir / name
        fig.tight_layout()
        fig.savefig(p)
        plt.close(fig)
        files.append(p)

    # 1. 시계열
    fig, ax1 = plt.subplots(figsize=(13, 5))
    ax2 = ax1.twinx()
    ax2.plot(pressure["time"], pressure["pressure"], color="0.55", lw=1, label="Pressure")
    ax2.set_ylabel("Atmospheric pressure (hPa)")
    ax2.spines["right"].set_visible(True)
    ax2.grid(False)
    for e in eqps:
        g = df[df["eqp"] == e]
        ax1.scatter(g["time"], g["thickness"], s=12, color=cmap[e], label=e, zorder=3)
    ax1.set_ylabel("SOH thickness (Å)")
    ax1.set_title("Pressure & SOH thickness over time")
    ax1.set_zorder(ax2.get_zorder() + 1)
    ax1.patch.set_visible(False)
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc="upper left", fontsize=8, ncol=len(eqps) + 1)
    save(fig, "1_timeseries.png")

    # 2. 산점도 (설비별)
    n = len(eqps)
    fig, axes = plt.subplots(1, n + 1, figsize=(4.2 * (n + 1), 4.2), squeeze=False)
    axes = axes[0]
    ax = axes[0]
    for e in eqps:
        g = df[df["eqp"] == e]
        ax.scatter(g["pressure"], g["thickness_dev"], s=10, color=cmap[e], label=e, alpha=0.7)
    _fit_line(ax, df["pressure"], df["thickness_dev"], "k")
    ax.set_title("All EQP (thickness − EQP mean)")
    ax.set_xlabel("Pressure (hPa)")
    ax.set_ylabel("Δ thickness (Å)")
    ax.legend(fontsize=8)
    for ax, e in zip(axes[1:], eqps):
        g = df[df["eqp"] == e]
        ax.scatter(g["pressure"], g["thickness"], s=10, color=cmap[e], alpha=0.7)
        _fit_line(ax, g["pressure"], g["thickness"], cmap[e])
        ax.set_title(e)
        ax.set_xlabel("Pressure (hPa)")
        ax.set_ylabel("Thickness (Å)")
    save(fig, "2_scatter.png")

    # 3. 대기압 구간별 박스플롯
    lo = np.floor(df["pressure"].min() / PRESSURE_BIN_HPA) * PRESSURE_BIN_HPA
    edges = np.arange(lo, df["pressure"].max() + PRESSURE_BIN_HPA, PRESSURE_BIN_HPA)
    b = pd.cut(df["pressure"], edges, right=False)
    groups = [(f"{iv.left:g}", g["thickness_dev"].to_numpy())
              for iv, g in df.groupby(b, observed=True) if len(g)]
    fig, ax = plt.subplots(figsize=(max(6, 0.7 * len(groups) + 2), 4.5))
    ax.boxplot([g for _, g in groups], showfliers=False)
    ax.set_xticks(range(1, len(groups) + 1), [l for l, _ in groups], rotation=45)
    for i, (_, g) in enumerate(groups, 1):
        ax.text(i, ax.get_ylim()[1], f"n={len(g)}", ha="center", va="bottom", fontsize=7)
    ax.set_xlabel(f"Pressure bin start (hPa, width {PRESSURE_BIN_HPA:g})")
    ax.set_ylabel("Δ thickness vs EQP mean (Å)")
    ax.set_title("Thickness by pressure bin", pad=16)
    save(fig, "3_boxplot_by_bin.png")

    # 4. 설비별 민감도
    s = summary[summary["구분"].str.startswith("설비 ")]
    fig, ax = plt.subplots(figsize=(max(5, 1.1 * len(s) + 2), 4))
    names = s["구분"].str.replace("설비 ", "", regex=False)
    ax.bar(names, s["기울기(Å/hPa)"], color=[cmap.get(e, "C0") for e in names])
    for i, (v, r2) in enumerate(zip(s["기울기(Å/hPa)"], s["R²"])):
        if pd.notna(v):
            ax.text(i, v, f"{v:.2f}\nR²={r2:.2f}", ha="center",
                    va="bottom" if v >= 0 else "top", fontsize=8)
    ax.axhline(0, color="k", lw=0.8)
    ax.set_ylabel("Slope (Å / hPa)")
    ax.set_title("Pressure sensitivity by EQP")
    save(fig, "4_sensitivity_by_eqp.png")

    # 5. 시차 분석
    fig, ax = plt.subplots(figsize=(7, 4))
    ax.plot(lag["시차(분, TKIN 이전)"], lag["Pearson r"], marker="o")
    best = lag.loc[lag["Pearson r"].abs().idxmax()] if lag["Pearson r"].notna().any() else None
    if best is not None:
        ax.axvline(best["시차(분, TKIN 이전)"], color="C3", ls="--",
                   label=f"max |r| at {best['시차(분, TKIN 이전)']:.0f} min")
        ax.legend()
    ax.set_xlabel("Lag: pressure taken N min before TKIN_TIME")
    ax.set_ylabel("Pearson r")
    ax.set_title("Lag analysis")
    save(fig, "5_lag.png")
    return files


def _fit_line(ax, x, y, color):
    m = x.notna() & y.notna()
    if m.sum() >= MIN_N_FOR_STATS and x[m].nunique() > 1:
        k, b = np.polyfit(x[m], y[m], 1)
        xs = np.linspace(x[m].min(), x[m].max(), 50)
        ax.plot(xs, k * xs + b, color=color, lw=1.5, label=f"{k:.2f} Å/hPa")
        ax.legend(fontsize=8)


# ============================================================
# HTML (인터랙티브 리포트, 외부 라이브러리/인터넷 없이 열림)
# ============================================================
def write_html(path, df, pressure, lag):
    import json

    def ms(t):
        return ((t - pd.Timestamp("1970-01-01")) // pd.Timedelta(milliseconds=1)).astype("int64").tolist()

    d = df.sort_values("time")
    p = pressure.sort_values("time")
    lag_rows = [{"lag": int(r["시차(분, TKIN 이전)"]), "n": int(r["N"]),
                 "r": None if pd.isna(r["Pearson r"]) else round(float(r["Pearson r"]), 4),
                 "slope": None if pd.isna(r["기울기(Å/hPa)"]) else round(float(r["기울기(Å/hPa)"]), 4)}
                for _, r in lag.iterrows()]
    t0, t1 = d["time"].min(), d["time"].max()
    sub = (f"TKIN {t0:%Y-%m-%d %H:%M} ~ {t1:%Y-%m-%d %H:%M} · 웨이퍼 {len(d)}장 · "
           f"설비 {d['eqp'].nunique()}대 · 대기압 매칭 허용 ±{MATCH_TOLERANCE_MIN}분 · "
           f"생성 {pd.Timestamp.now():%Y-%m-%d %H:%M}")
    data = {
        "eqps": sorted(d["eqp"].astype(str).unique().tolist()),
        "w": {"t": ms(d["time"]), "eqp": d["eqp"].astype(str).tolist(),
              "lot": d["lot"].astype(str).tolist(), "wafer": d["wafer"].astype(str).tolist(),
              "thk": d["thickness"].astype(float).round(2).tolist(),
              "p": d["pressure"].astype(float).round(3).tolist()},
        "p": {"t": ms(p["time"]), "p": p["pressure"].astype(float).round(3).tolist()},
        "lag": lag_rows,
        "meta": {"sub": sub, "bin": PRESSURE_BIN_HPA},
    }
    js = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    Path(path).write_text(HTML_TEMPLATE.replace("__DATA__", js), encoding="utf-8")


HTML_TEMPLATE = r'''<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>대기압 vs HSOH 두께</title>
<style>
:root {
  color-scheme: light;
  --page: #f9f9f7; --surface: #fcfcfb; --ink: #0b0b0b; --ink2: #52514e; --muted: #898781;
  --grid: #e1e0d9; --axis: #c3c2b7; --border: rgba(11,11,11,0.10);
  --s1: #2a78d6; --s2: #eb6834; --s3: #1baf7a; --s4: #eda100;
  --s5: #e87ba4; --s6: #008300; --s7: #6250d6; --s8: #e34948; --s0: #898781;
  --sel: rgba(42,120,214,0.12);
}
@media (prefers-color-scheme: dark) {
  :root:not([data-theme="light"]) {
    color-scheme: dark;
    --page: #0d0d0d; --surface: #1a1a19; --ink: #f0efec; --ink2: #c3c2b7; --muted: #898781;
    --grid: #2c2c2a; --axis: #383835; --border: rgba(255,255,255,0.10);
    --s1: #3987e5; --s2: #d95926; --s3: #199e70; --s4: #c98500;
    --s5: #d55181; --s6: #008300; --s7: #9085e9; --s8: #e66767;
    --sel: rgba(57,135,229,0.18);
  }
}
:root[data-theme="dark"] {
  color-scheme: dark;
  --page: #0d0d0d; --surface: #1a1a19; --ink: #f0efec; --ink2: #c3c2b7; --muted: #898781;
  --grid: #2c2c2a; --axis: #383835; --border: rgba(255,255,255,0.10);
  --s1: #3987e5; --s2: #d95926; --s3: #199e70; --s4: #c98500;
  --s5: #d55181; --s6: #008300; --s7: #9085e9; --s8: #e66767;
  --sel: rgba(57,135,229,0.18);
}
* { box-sizing: border-box; }
body { margin: 0; background: var(--page); color: var(--ink);
  font: 14px/1.45 "Malgun Gothic", "Apple SD Gothic Neo", "Noto Sans KR", system-ui, sans-serif; }
.wrap { max-width: 1280px; margin: 0 auto; padding: 20px 16px 40px; }
h1 { font-size: 20px; margin: 0 0 2px; }
h2 { font-size: 15px; margin: 0; }
.sub, .hint { color: var(--ink2); font-size: 12px; margin: 2px 0 0; }
.filters { display: flex; flex-wrap: wrap; gap: 8px; align-items: center; margin: 16px 0 12px; }
.filters .lab { color: var(--ink2); font-size: 12px; margin-right: 2px; }
.chip { display: inline-flex; align-items: center; gap: 6px; padding: 4px 10px; border-radius: 999px;
  border: 1px solid var(--border); background: var(--surface); color: var(--ink); cursor: pointer; font: inherit; font-size: 13px; }
.chip .key { width: 14px; height: 3px; border-radius: 2px; }
.chip[aria-pressed="false"] { color: var(--muted); }
.chip[aria-pressed="false"] .key { opacity: 0.25; }
.btn { padding: 4px 10px; border-radius: 6px; border: 1px solid var(--border); background: var(--surface);
  color: var(--ink); cursor: pointer; font: inherit; font-size: 12px; }
.range { color: var(--ink2); font-size: 12px; }
.kpis { display: grid; grid-template-columns: repeat(auto-fit, minmax(170px, 1fr)); gap: 12px; margin-bottom: 12px; }
.kpi { background: var(--surface); border: 1px solid var(--border); border-radius: 10px; padding: 12px 14px; }
.kpi .v { font-size: 22px; font-weight: 700; font-variant-numeric: tabular-nums; }
.kpi .l { color: var(--ink2); font-size: 12px; }
.card { background: var(--surface); border: 1px solid var(--border); border-radius: 10px; padding: 14px; margin-bottom: 12px; min-width: 0; }
.head { display: flex; justify-content: space-between; align-items: baseline; gap: 8px; flex-wrap: wrap; margin-bottom: 6px; }
.grid2 { display: grid; grid-template-columns: 1fr 1fr; gap: 12px; }
@media (max-width: 860px) { .grid2 { grid-template-columns: 1fr; } }
svg { display: block; width: 100%; overflow: visible; user-select: none; }
svg text { fill: var(--muted); font-size: 11px; font-variant-numeric: tabular-nums; }
svg .grid { stroke: var(--grid); stroke-width: 1; }
svg .base { stroke: var(--axis); stroke-width: 1; }
svg .lbl { fill: var(--ink2); font-size: 12px; }
svg .val { fill: var(--ink); font-size: 11px; }
.seg { display: inline-flex; border: 1px solid var(--border); border-radius: 6px; overflow: hidden; }
.seg button { border: 0; background: var(--surface); color: var(--ink2); padding: 3px 10px; font: inherit; font-size: 12px; cursor: pointer; }
.seg button[aria-pressed="true"] { background: var(--sel); color: var(--ink); font-weight: 600; }
.tip { position: fixed; pointer-events: none; z-index: 10; background: var(--surface); color: var(--ink);
  border: 1px solid var(--border); border-radius: 8px; padding: 8px 10px; font-size: 12px;
  box-shadow: 0 4px 16px rgba(0,0,0,0.12); display: none; min-width: 150px; }
.tip .t { color: var(--ink2); margin-bottom: 4px; }
.tip .r { display: flex; align-items: center; gap: 6px; }
.tip .r b { font-variant-numeric: tabular-nums; }
.tip .r span.k { width: 12px; height: 2px; border-radius: 1px; display: inline-block; }
.tip .r span.n { color: var(--ink2); }
table { width: 100%; border-collapse: collapse; font-size: 13px; font-variant-numeric: tabular-nums; }
th, td { text-align: right; padding: 6px 8px; border-bottom: 1px solid var(--grid); }
th:first-child, td:first-child { text-align: left; }
th { color: var(--ink2); font-weight: 600; }
.note { color: var(--muted); font-size: 12px; margin-top: 6px; }
.axbox { display: flex; flex-wrap: wrap; gap: 10px 18px; align-items: center; padding: 10px 12px; margin: 0 0 10px;
  border: 1px solid var(--border); border-radius: 8px; background: var(--page); font-size: 12px; color: var(--ink2); }
.axbox[hidden] { display: none; }
.axbox .grp { display: inline-flex; align-items: center; gap: 6px; flex-wrap: wrap; }
.axbox .grp b { color: var(--ink); font-weight: 600; }
.axbox label { display: inline-flex; align-items: center; gap: 4px; }
.axbox input, .axbox select { padding: 3px 6px; border: 1px solid var(--border); border-radius: 4px;
  background: var(--surface); color: var(--ink); font: inherit; }
.axbox input[type=number] { width: 104px; }
.axbox input[type=datetime-local] { width: 190px; }
.legend { display: flex; flex-wrap: wrap; gap: 4px 14px; font-size: 12px; color: var(--ink2); }
.legend .it { display: inline-flex; align-items: center; gap: 6px; cursor: pointer; background: none; border: 0;
  padding: 2px 0; color: inherit; font: inherit; }
.legend .it[aria-pressed="false"] { opacity: 0.35; text-decoration: line-through; }
.legend .ln { width: 16px; height: 2px; border-radius: 1px; }
.legend .dt { width: 8px; height: 8px; border-radius: 50%; }
.tsbox { display: flex; gap: 10px; }
.tsbox > #ts { flex: 1; min-width: 0; }
.tsbox.pos-top { flex-direction: column; }
.tsbox.pos-bottom { flex-direction: column-reverse; }
.tsbox.pos-right { flex-direction: row-reverse; }
.tsbox.pos-right .legend { flex-direction: column; flex-wrap: nowrap; min-width: 110px; padding-top: 20px; }
.tsbox.pos-none .legend { display: none; }
</style>
</head>
<body>
<div class="wrap">
  <h1>대기압 vs HSOH 두께</h1>
  <p class="sub" id="sub"></p>

  <div class="filters">
    <span class="lab">설비</span>
    <span id="eqpChips" style="display:contents"></span>
    <button class="btn" id="allOn">전체 선택</button>
    <button class="btn" id="allOff">전체 해제</button>
    <span class="range" id="rangeLabel"></span>
    <button class="btn" id="rangeReset" hidden>기간 해제</button>
  </div>

  <div class="kpis" id="kpis"></div>

  <section class="card">
    <div class="head"><h2>시간에 따른 대기압과 두께</h2>
      <span class="hint">아래 그래프에서 좌우로 드래그하면 기간을 선택합니다 (더블클릭: 해제)</span></div>
    <div class="tsbox pos-top" id="tsbox"><div class="legend" id="tsLegend"></div><div id="ts"></div></div>
  </section>

  <div class="grid2">
    <section class="card">
      <div class="head"><h2>대기압 vs 두께</h2>
        <div class="seg" id="yMode"><button data-v="dev" aria-pressed="true">설비 평균 대비 편차</button><button data-v="raw" aria-pressed="false">두께 원값</button></div></div>
      <div id="sc"></div>
    </section>
    <section class="card">
      <div class="head"><h2>설비별 민감도 (Å/hPa)</h2><span class="hint">선택한 기간·설비 기준</span></div>
      <div id="sens"></div>
    </section>
  </div>

  <div class="grid2">
    <section class="card">
      <div class="head"><h2>대기압 구간별 두께 편차</h2><span class="hint">막대: 평균, 선: ±1σ</span></div>
      <div id="bins"></div>
    </section>
    <section class="card">
      <div class="head"><h2>시차 분석</h2><span class="hint">TKIN_TIME 보다 N분 앞선 대기압과의 상관 (전체 데이터 기준)</span></div>
      <div id="lag"></div>
    </section>
  </div>

  <section class="card">
    <div class="head"><h2>요약 표</h2><span class="hint">선택한 기간·설비 기준</span></div>
    <div id="tbl" style="overflow-x:auto"></div>
    <p class="note">p-value 는 엑셀 리포트(요약_회귀 시트, 전체 데이터 기준)를 참고하세요.</p>
  </section>
</div>
<div class="tip" id="tip"></div>

<script>
const D = __DATA__;
const NS = "http://www.w3.org/2000/svg";
const N = D.w.t.length;
const EQPS = D.eqps;
const colorOf = e => { const i = EQPS.indexOf(e); return i >= 0 && i < 8 ? `var(--s${i + 1})` : "var(--s0)"; };
const state = { on: new Set(EQPS), range: null, y: "dev" };

// ---------- 공통 유틸 ----------
const el = (tag, attrs = {}, parent) => {
  const n = document.createElementNS(NS, tag);
  for (const [k, v] of Object.entries(attrs)) {
    if (k === "text") n.textContent = v; else if (k === "style") n.style.cssText = v; else n.setAttribute(k, v);
  }
  if (parent) parent.appendChild(n);
  return n;
};
const lin = (d0, d1, r0, r1) => { const k = d1 === d0 ? 0 : (r1 - r0) / (d1 - d0); const f = v => r0 + (v - d0) * k; f.inv = p => d0 + (p - r0) / (k || 1); return f; };
function niceTicks(lo, hi, n = 5) {
  if (!(hi > lo)) { const c = isFinite(lo) ? lo : 0; return [c]; }
  const raw = (hi - lo) / n, mag = Math.pow(10, Math.floor(Math.log10(raw)));
  const step = [1, 2, 2.5, 5, 10].map(m => m * mag).find(s => (hi - lo) / s <= n) || 10 * mag;
  const out = []; for (let v = Math.ceil(lo / step) * step; v <= hi + 1e-9; v += step) out.push(+v.toFixed(10));
  return out;
}
const pad2 = v => String(v).padStart(2, "0");
const fmtDT = ms => { const d = new Date(ms); return `${d.getUTCFullYear()}-${pad2(d.getUTCMonth() + 1)}-${pad2(d.getUTCDate())} ${pad2(d.getUTCHours())}:${pad2(d.getUTCMinutes())}`; };
const fmtD = ms => { const d = new Date(ms); return `${pad2(d.getUTCMonth() + 1)}/${pad2(d.getUTCDate())}`; };
const fmtH = ms => { const d = new Date(ms); return `${pad2(d.getUTCHours())}:${pad2(d.getUTCMinutes())}`; };
function timeTicks(t0, t1, n) {
  const H = 3600e3, Dd = 24 * H, span = t1 - t0;
  const steps = [H, 2 * H, 3 * H, 6 * H, 12 * H, Dd, 2 * Dd, 7 * Dd, 14 * Dd, 30 * Dd];
  const st = steps.find(s => span / s <= n) || 30 * Dd;
  const out = []; for (let v = Math.ceil(t0 / st) * st; v <= t1; v += st) out.push(v);
  return { ticks: out, fmt: st >= Dd ? fmtD : (v => (new Date(v).getUTCHours() === 0 ? fmtD(v) : fmtH(v))) };
}
const f = (v, d = 2) => (v == null || !isFinite(v)) ? "-" : v.toFixed(d);
const tk = v => String(v === 0 ? 0 : v);  // 눈금 값은 niceTicks 가 이미 깔끔하게 만든 수
function regress(xs, ys) {
  const n = xs.length; if (n < 3) return { n, slope: NaN, icpt: NaN, r: NaN, r2: NaN };
  let mx = 0, my = 0; for (let i = 0; i < n; i++) { mx += xs[i]; my += ys[i]; } mx /= n; my /= n;
  let sxx = 0, syy = 0, sxy = 0;
  for (let i = 0; i < n; i++) { const dx = xs[i] - mx, dy = ys[i] - my; sxx += dx * dx; syy += dy * dy; sxy += dx * dy; }
  if (sxx === 0) return { n, slope: NaN, icpt: NaN, r: NaN, r2: NaN };
  const slope = sxy / sxx, r = syy === 0 ? 0 : sxy / Math.sqrt(sxx * syy);
  return { n, slope, icpt: my - slope * mx, r, r2: r * r, mean: my, sd: Math.sqrt(syy / Math.max(n - 1, 1)) };
}

// ---------- 툴팁 ----------
const tip = document.getElementById("tip");
function showTip(ev, title, rows) {
  tip.replaceChildren();
  const t = document.createElement("div"); t.className = "t"; t.textContent = title; tip.appendChild(t);
  for (const [color, val, name] of rows) {
    const r = document.createElement("div"); r.className = "r";
    if (color) { const k = document.createElement("span"); k.className = "k"; k.style.background = color; r.appendChild(k); }
    const b = document.createElement("b"); b.textContent = val; r.appendChild(b);
    const s = document.createElement("span"); s.className = "n"; s.textContent = name; r.appendChild(s);
    tip.appendChild(r);
  }
  tip.style.display = "block";
  const w = tip.offsetWidth, h = tip.offsetHeight;
  let x = ev.clientX + 14, y = ev.clientY + 14;
  if (x + w > innerWidth - 8) x = ev.clientX - w - 14;
  if (y + h > innerHeight - 8) y = ev.clientY - h - 14;
  tip.style.left = x + "px"; tip.style.top = y + "px";
}
const hideTip = () => { tip.style.display = "none"; };

// ---------- 데이터 선택 ----------
function selected() {
  const idx = [];
  for (let i = 0; i < N; i++) {
    if (!state.on.has(D.w.eqp[i])) continue;
    if (state.range && (D.w.t[i] < state.range[0] || D.w.t[i] > state.range[1])) continue;
    idx.push(i);
  }
  const sum = {}, cnt = {};
  for (const i of idx) { const e = D.w.eqp[i]; sum[e] = (sum[e] || 0) + D.w.thk[i]; cnt[e] = (cnt[e] || 0) + 1; }
  const dev = new Map(); for (const i of idx) { const e = D.w.eqp[i]; dev.set(i, D.w.thk[i] - sum[e] / cnt[e]); }
  return { idx, dev };
}

// ---------- 필터 UI ----------
const chips = document.getElementById("eqpChips");
for (const e of EQPS) {
  const b = document.createElement("button"); b.className = "chip"; b.setAttribute("aria-pressed", "true");
  const k = document.createElement("span"); k.className = "key"; k.style.background = colorOf(e); b.appendChild(k);
  b.appendChild(document.createTextNode(e));
  b.onclick = () => { state.on.has(e) ? state.on.delete(e) : state.on.add(e); b.setAttribute("aria-pressed", state.on.has(e)); render(); };
  chips.appendChild(b);
}
const setAll = on => { state.on = new Set(on ? EQPS : []); chips.querySelectorAll(".chip").forEach(c => c.setAttribute("aria-pressed", on)); render(); };
document.getElementById("allOn").onclick = () => setAll(true);
document.getElementById("allOff").onclick = () => setAll(false);
document.getElementById("rangeReset").onclick = () => { state.range = null; render(); };
document.querySelectorAll("#yMode button").forEach(b => b.onclick = () => {
  state.y = b.dataset.v; document.querySelectorAll("#yMode button").forEach(x => x.setAttribute("aria-pressed", x === b)); render();
});
document.getElementById("sub").textContent = D.meta.sub;

// ---------- KPI ----------
function renderKpis(S) {
  const xs = S.idx.map(i => D.w.p[i]), ys = S.idx.map(i => S.dev.get(i));
  const r = regress(xs, ys);
  const pmin = Math.min(...xs), pmax = Math.max(...xs);
  const items = [
    [S.idx.length.toLocaleString(), "웨이퍼 수"],
    [f(r.slope, 2) + " Å/hPa", "대기압 1 hPa 당 두께 변화 (설비 평균 제거)"],
    [f(r.r2, 3), "R²"],
    [xs.length ? `${f(pmin, 1)} ~ ${f(pmax, 1)}` : "-", "대기압 범위 (hPa)"],
  ];
  const box = document.getElementById("kpis"); box.replaceChildren();
  for (const [v, l] of items) {
    const d = document.createElement("div"); d.className = "kpi";
    const a = document.createElement("div"); a.className = "v"; a.textContent = v;
    const b = document.createElement("div"); b.className = "l"; b.textContent = l;
    d.append(a, b); box.appendChild(d);
  }
}

// ---------- 축 설정 / 범례 설정 ----------
// AX[차트][축] = [라벨, 종류]. 입력칸이 비어 있으면 자동 범위.
const AX = {
  ts: { t: ["시간", "time"], p: ["대기압 (hPa)"], y: ["두께 (Å)"] },
  sc: { x: ["대기압 (hPa)"], y: ["Y축 (Å)"] },
  sens: { x: ["기울기 (Å/hPa)"] },
  bins: { y: ["Δ 두께 (Å)"] },
  lag: { y: ["Pearson r"] },
};
state.ax = {}; state.auto = {}; state.legendPos = "top"; state.tsHide = new Set(); state.tsLayout = "overlay";
const toInputTime = ms => new Date(ms).toISOString().slice(0, 16);
const fromInputTime = v => (v ? Date.parse(v + ":00Z") : null);
function rng(c, a, lo, hi) {
  state.auto[c + "." + a] = [lo, hi];
  const o = (state.ax[c] && state.ax[c][a]) || [null, null];
  const L = o[0] != null ? o[0] : lo, H = o[1] != null ? o[1] : hi;
  return H > L ? [L, H] : [lo, hi];
}
let clipN = 0;
function clipG(svg, x, y, w, h) {
  const id = "clip" + (++clipN), cp = el("clipPath", { id }, el("defs", {}, svg));
  el("rect", { x, y, width: Math.max(0, w), height: Math.max(0, h) }, cp);
  return el("g", { "clip-path": `url(#${id})` }, svg);
}
// 점이 많으면 SVG 요소 대신 canvas 에 한 번에 그린다 (수만 개도 빠름). svg 는 그 위에 겹쳐 축/선/툴팁 담당
function dotLayer(box, svg, W, H, clip) {
  box.style.position = "relative"; svg.style.position = "relative";
  const c = document.createElement("canvas"), dpr = window.devicePixelRatio || 1;
  c.width = Math.round(W * dpr); c.height = Math.round(H * dpr);
  c.style.cssText = `position:absolute;left:0;top:0;width:${W}px;height:${H}px;pointer-events:none`;
  box.insertBefore(c, svg);
  const ctx = c.getContext("2d"); ctx.scale(dpr, dpr);
  ctx.beginPath(); ctx.rect(clip[0], clip[1], clip[2], clip[3]); ctx.clip();
  return ctx;
}
function cssVar(v) {
  const mm = /var\((--[\w-]+)\)/.exec(v);
  return mm ? getComputedStyle(document.documentElement).getPropertyValue(mm[1]).trim() : v;
}
function drawDots(ctx, idx, px, py, r, alpha) {
  const groups = new Map();
  for (const i of idx) { const e = D.w.eqp[i]; if (!groups.has(e)) groups.set(e, []); groups.get(e).push(i); }
  const ring = cssVar("var(--surface)");
  ctx.globalAlpha = alpha;
  for (const [e, ii] of groups) {
    ctx.fillStyle = cssVar(colorOf(e)); ctx.strokeStyle = ring; ctx.lineWidth = 1;
    ctx.beginPath();
    for (const i of ii) { const x = px(i), y = py(i); ctx.moveTo(x + r, y); ctx.arc(x, y, r, 0, 6.2832); }
    if (ii.length < 3000) ctx.stroke();
    ctx.fill();
  }
  ctx.globalAlpha = 1;
}
for (const [cid, axes] of Object.entries(AX)) {
  const card = document.getElementById(cid).closest(".card"), head = card.querySelector(".head");
  const btn = document.createElement("button"); btn.className = "btn"; btn.textContent = "⚙ 축 설정";
  btn.setAttribute("aria-expanded", "false"); head.appendChild(btn);
  const box = document.createElement("div"); box.className = "axbox"; box.hidden = true; head.after(box);
  btn.onclick = () => { box.hidden = !box.hidden; btn.setAttribute("aria-expanded", String(!box.hidden)); };
  state.ax[cid] = {};
  for (const [a, [label, kind]] of Object.entries(axes)) {
    state.ax[cid][a] = [null, null];
    const g = document.createElement("span"); g.className = "grp";
    const b = document.createElement("b"); b.textContent = label; g.appendChild(b);
    ["최소", "최대"].forEach((nm, j) => {
      const lab = document.createElement("label"); lab.append(nm);
      const inp = document.createElement("input");
      inp.type = kind === "time" ? "datetime-local" : "number";
      if (kind !== "time") inp.step = "any";
      inp.dataset.key = cid + "." + a; inp.dataset.j = j; inp.dataset.kind = kind || "num";
      inp.addEventListener("change", () => {
        const v = inp.value === "" ? null : (kind === "time" ? fromInputTime(inp.value) : Number(inp.value));
        state.ax[cid][a][j] = v == null || isNaN(v) ? null : v; render();
      });
      lab.appendChild(inp); g.appendChild(lab);
    });
    box.appendChild(g);
  }
  if (cid === "ts") {
    const g = document.createElement("span"); g.className = "grp";
    const b = document.createElement("b"); b.textContent = "범례"; g.appendChild(b);
    const sel = document.createElement("select");
    for (const [v, t] of [["top", "위"], ["bottom", "아래"], ["right", "오른쪽"], ["none", "숨김"]]) {
      const o = document.createElement("option"); o.value = v; o.textContent = t; sel.appendChild(o);
    }
    sel.onchange = () => { state.legendPos = sel.value; render(); };
    g.appendChild(sel); box.appendChild(g);
    const g2 = document.createElement("span"); g2.className = "grp";
    const b2 = document.createElement("b"); b2.textContent = "배치"; g2.appendChild(b2);
    const sel2 = document.createElement("select");
    for (const [v, t] of [["overlay", "겹치기 (왼쪽 두께 / 오른쪽 대기압)"], ["split", "위아래 분리"]]) {
      const o = document.createElement("option"); o.value = v; o.textContent = t; sel2.appendChild(o);
    }
    sel2.onchange = () => { state.tsLayout = sel2.value; render(); };
    g2.appendChild(sel2); box.appendChild(g2);
  }
  const rs = document.createElement("button"); rs.className = "btn"; rs.textContent = "자동으로";
  rs.onclick = () => {
    box.querySelectorAll("input").forEach(i => { i.value = ""; });
    for (const a in state.ax[cid]) state.ax[cid][a] = [null, null];
    render();
  };
  box.appendChild(rs);
}
function updatePlaceholders() {
  document.querySelectorAll(".axbox input").forEach(inp => {
    const au = state.auto[inp.dataset.key]; if (!au) return;
    const v = au[+inp.dataset.j];
    inp.placeholder = "자동 " + (inp.dataset.kind === "time" ? fmtDT(v) : (Math.abs(v) >= 100 ? f(v, 1) : f(v, 3)));
  });
}

// ---------- 1. 시계열 (겹치기: 왼쪽 두께/오른쪽 대기압, 또는 위아래 분리) ----------
let legendSig = "";
function renderLegend() {
  // 버튼을 매번 새로 만들면 입력칸 change 직후의 클릭이 사라지므로, 항목이 바뀔 때만 다시 만든다
  const lg = document.getElementById("tsLegend");
  document.getElementById("tsbox").className = "tsbox pos-" + state.legendPos;
  const items = [["__p", "대기압", "var(--ink2)", "ln"]].concat(EQPS.filter(e => state.on.has(e)).map(e => [e, e, colorOf(e), "dt"]));
  const sig = items.map(it => it[0]).join("|");
  if (sig !== legendSig) {
    legendSig = sig; lg.replaceChildren();
    for (const [key, name, color, shape] of items) {
      const b = document.createElement("button"); b.className = "it"; b.dataset.key = key;
      b.title = "클릭: 이 차트에서 숨기기/보이기";
      const k = document.createElement("span"); k.className = shape; k.style.background = color; b.appendChild(k);
      b.appendChild(document.createTextNode(name));
      b.onclick = () => { state.tsHide.has(key) ? state.tsHide.delete(key) : state.tsHide.add(key); render(); };
      lg.appendChild(b);
    }
  }
  lg.querySelectorAll(".it").forEach(b => b.setAttribute("aria-pressed", String(!state.tsHide.has(b.dataset.key))));
}
function renderTS(S) {
  renderLegend();
  const box = document.getElementById("ts"); box.replaceChildren();
  const ov = state.tsLayout === "overlay";  // 겹치기: 왼쪽 두께 / 오른쪽 대기압
  const W = box.clientWidth || 900, m = { l: 56, r: ov ? 60 : 12 }, pw = W - m.l - m.r;
  let pT, pBt, tT, tBt;
  if (ov) { pT = tT = 18; pBt = tBt = 18 + 340; } else { pT = 8; pBt = 138; tT = 164; tBt = 384; }
  const H = tBt + 26, top = Math.min(pT, tT);
  const svg = el("svg", { viewBox: `0 0 ${W} ${H}`, height: H }, box);
  const showP = !state.tsHide.has("__p");
  const vis = S.idx.filter(i => !state.tsHide.has(D.w.eqp[i]));
  const t0 = Math.min(D.p.t[0], N ? D.w.t[0] : Infinity), t1 = Math.max(D.p.t[D.p.t.length - 1], N ? D.w.t[N - 1] : -Infinity);
  const [tA, tB] = rng("ts", "t", t0, t1), x = lin(tA, tB, m.l, W - m.r);
  const pv = D.p.p, plo = Math.min(...pv), phi = Math.max(...pv), pp = (phi - plo) * 0.08 || 1;
  const [pA, pB] = rng("ts", "p", plo - pp, phi + pp), y1 = lin(pA, pB, pBt, pT);
  const tv = vis.map(i => D.w.thk[i]);
  const tlo = tv.length ? Math.min(...tv) : 0, thi = tv.length ? Math.max(...tv) : 1, tp = (thi - tlo) * 0.08 || 1;
  const [yA, yB] = rng("ts", "y", tlo - tp, thi + tp), y2 = lin(yA, yB, tBt, tT);
  const tt = timeTicks(tA, tB, Math.max(4, Math.floor(W / 110)));
  // 축/그리드 (겹치기일 때 그리드는 왼쪽 두께 축 기준으로만)
  for (const v of niceTicks(yA, yB, 5)) { el("line", { class: "grid", x1: m.l, x2: W - m.r, y1: y2(v), y2: y2(v) }, svg); el("text", { x: m.l - 6, y: y2(v) + 4, "text-anchor": "end", text: tk(v) }, svg); }
  for (const v of niceTicks(pA, pB, ov ? 5 : 4)) {
    if (ov) el("text", { x: W - m.r + 6, y: y1(v) + 4, "text-anchor": "start", text: f(v, 1) }, svg);
    else { el("line", { class: "grid", x1: m.l, x2: W - m.r, y1: y1(v), y2: y1(v) }, svg); el("text", { x: m.l - 6, y: y1(v) + 4, "text-anchor": "end", text: f(v, 1) }, svg); }
  }
  for (const v of tt.ticks) el("text", { x: x(v), y: tBt + 16, "text-anchor": "middle", text: tt.fmt(v) }, svg);
  el("line", { class: "base", x1: m.l, x2: W - m.r, y1: tBt, y2: tBt }, svg);
  if (ov) {
    el("line", { class: "base", x1: m.l, x2: m.l, y1: tT, y2: tBt }, svg);
    el("line", { class: "base", x1: W - m.r, x2: W - m.r, y1: pT, y2: pBt }, svg);
    el("text", { class: "lbl", x: m.l, y: tT - 6, "text-anchor": "end", text: "두께 (Å)" }, svg);
    el("text", { class: "lbl", x: W - m.r, y: pT - 6, "text-anchor": "start", text: "대기압 (hPa)" }, svg);
  } else {
    el("line", { class: "base", x1: m.l, x2: W - m.r, y1: pBt, y2: pBt }, svg);
    el("text", { class: "lbl", x: m.l, y: pT, dy: "-2", text: "대기압 (hPa)" }, svg);
    el("text", { class: "lbl", x: m.l, y: tT - 4, text: "두께 (Å)" }, svg);
  }
  // 선택 기간
  if (state.range) {
    const gs = clipG(svg, m.l, top, pw, tBt - top);
    el("rect", { x: x(state.range[0]), y: top, width: Math.max(1, x(state.range[1]) - x(state.range[0])), height: tBt - top, style: "fill:var(--sel)" }, gs);
  }
  // 대기압 선
  if (showP) {
    const g1 = clipG(svg, m.l, pT, pw, pBt - pT);
    let d = ""; for (let i = 0; i < D.p.t.length; i++) d += (i ? "L" : "M") + x(D.p.t[i]).toFixed(1) + "," + y1(pv[i]).toFixed(1);
    el("path", { d, style: "fill:none;stroke:var(--ink2);stroke-width:2;stroke-linejoin:round" }, g1);
  }
  // 두께 점 (canvas)
  drawDots(dotLayer(box, svg, W, H, [m.l, tT, pw, tBt - tT]), vis, i => x(D.w.t[i]), i => y2(D.w.thk[i]), 3, 1);
  // 크로스헤어 + 브러시
  const cross = el("line", { x1: 0, x2: 0, y1: top, y2: tBt, style: "stroke:var(--muted);stroke-width:1;display:none" }, svg);
  const dot = el("circle", { r: 4, style: "fill:var(--ink2);stroke:var(--surface);stroke-width:2;display:none" }, svg);
  const ring = el("circle", { r: 6, style: "fill:none;stroke:var(--ink);stroke-width:1.5;display:none" }, svg);
  const brush = el("rect", { y: top, height: tBt - top, width: 0, style: "fill:var(--sel);stroke:var(--s1);stroke-width:1;display:none" }, svg);
  const hit = el("rect", { x: m.l, y: top, width: pw, height: tBt - top, style: "fill:transparent;cursor:crosshair" }, svg);
  const toX = ev => { const r = svg.getBoundingClientRect(); return (ev.clientX - r.left) * W / r.width; };
  const toY = ev => { const r = svg.getBoundingClientRect(); return (ev.clientY - r.top) * H / r.height; };
  const nearestWafer = (px, py) => {
    let best = -1, bd = 400;
    for (const i of vis) { const dx = x(D.w.t[i]) - px, dy = y2(D.w.thk[i]) - py, dd = dx * dx + dy * dy; if (dd < bd) { bd = dd; best = i; } }
    return best;
  };
  const showWafer = (ev, best) => {
    cross.style.display = "none"; dot.style.display = "none";
    ring.style.display = ""; ring.setAttribute("cx", x(D.w.t[best])); ring.setAttribute("cy", y2(D.w.thk[best]));
    showTip(ev, fmtDT(D.w.t[best]) + " (TKIN)", [
      [colorOf(D.w.eqp[best]), f(D.w.thk[best], 1) + " Å", "두께"],
      [null, f(D.w.p[best], 2) + " hPa", "대기압"],
      [null, D.w.eqp[best], "설비"], [null, D.w.lot[best], "LOT"], [null, D.w.wafer[best], "WAFER"]]);
  };
  let drag = null;
  hit.addEventListener("pointerdown", ev => { drag = toX(ev); hit.setPointerCapture(ev.pointerId); });
  hit.addEventListener("pointermove", ev => {
    const px = Math.min(Math.max(toX(ev), m.l), W - m.r), py = toY(ev);
    if (drag != null) {
      brush.style.display = ""; brush.setAttribute("x", Math.min(drag, px)); brush.setAttribute("width", Math.abs(px - drag));
      hideTip(); return;
    }
    if (ov || py > tT - 6) { // 웨이퍼 점 근처면 웨이퍼 정보
      const best = nearestWafer(px, py);
      if (best >= 0) { showWafer(ev, best); return; }
      if (!ov) { cross.style.display = dot.style.display = ring.style.display = "none"; hideTip(); return; }
    }
    ring.style.display = "none";
    if (!showP) { cross.style.display = dot.style.display = "none"; hideTip(); return; }
    const tq = x.inv(px);
    let lo = 0, hi = D.p.t.length - 1;
    while (hi - lo > 1) { const mid = (lo + hi) >> 1; if (D.p.t[mid] < tq) lo = mid; else hi = mid; }
    const k = Math.abs(D.p.t[lo] - tq) < Math.abs(D.p.t[hi] - tq) ? lo : hi;
    const cx = x(D.p.t[k]);
    cross.style.display = ""; cross.setAttribute("x1", cx); cross.setAttribute("x2", cx);
    dot.style.display = ""; dot.setAttribute("cx", cx); dot.setAttribute("cy", y1(pv[k]));
    showTip(ev, fmtDT(D.p.t[k]), [["var(--ink2)", f(pv[k], 2) + " hPa", "대기압"]]);
  });
  hit.addEventListener("pointerup", ev => {
    if (drag == null) return;
    const px = Math.min(Math.max(toX(ev), m.l), W - m.r), a = Math.min(drag, px), b = Math.max(drag, px);
    drag = null; brush.style.display = "none";
    if (b - a > 6) { state.range = [x.inv(a), x.inv(b)]; render(); }
  });
  hit.addEventListener("dblclick", () => { state.range = null; render(); });
  hit.addEventListener("pointerleave", () => { if (drag == null) { cross.style.display = dot.style.display = ring.style.display = "none"; hideTip(); } });
}

// ---------- 2. 산점도 ----------
function renderScatter(S) {
  const box = document.getElementById("sc"); box.replaceChildren();
  const W = box.clientWidth || 500, H = 340, m = { l: 52, r: 70, t: 22, b: 36 };
  const svg = el("svg", { viewBox: `0 0 ${W} ${H}`, height: H }, box);
  const yv = i => state.y === "dev" ? S.dev.get(i) : D.w.thk[i];
  if (!S.idx.length) { el("text", { x: W / 2, y: H / 2, "text-anchor": "middle", text: "선택된 데이터가 없습니다" }, svg); return; }
  const xs = S.idx.map(i => D.w.p[i]), ys = S.idx.map(yv);
  const xlo = Math.min(...xs), xhi = Math.max(...xs), xp = (xhi - xlo) * 0.05 || 0.5;
  const ylo = Math.min(...ys), yhi = Math.max(...ys), yp = (yhi - ylo) * 0.08 || 1;
  const [xA, xB] = rng("sc", "x", xlo - xp, xhi + xp), [yA, yB] = rng("sc", "y", ylo - yp, yhi + yp);
  const x = lin(xA, xB, m.l, W - m.r), y = lin(yA, yB, H - m.b, m.t);
  for (const v of niceTicks(yA, yB, 5)) { el("line", { class: "grid", x1: m.l, x2: W - m.r, y1: y(v), y2: y(v) }, svg); el("text", { x: m.l - 6, y: y(v) + 4, "text-anchor": "end", text: tk(v) }, svg); }
  for (const v of niceTicks(xA, xB, 6)) el("text", { x: x(v), y: H - m.b + 16, "text-anchor": "middle", text: tk(v) }, svg);
  el("line", { class: "base", x1: m.l, x2: W - m.r, y1: H - m.b, y2: H - m.b }, svg);
  el("text", { class: "lbl", x: (m.l + W - m.r) / 2, y: H - 4, "text-anchor": "middle", text: "대기압 (hPa)" }, svg);
  el("text", { class: "lbl", x: 4, y: m.t - 2, text: state.y === "dev" ? "Δ 두께 (Å)" : "두께 (Å)" }, svg);
  const g = clipG(svg, m.l, m.t, W - m.l - m.r, H - m.t - m.b);
  drawDots(dotLayer(box, svg, W, H, [m.l, m.t, W - m.l - m.r, H - m.t - m.b]), S.idx, i => x(D.w.p[i]), i => y(yv(i)), 3.5, 0.75);
  // 설비별 회귀선 + 직접 라벨
  const on = EQPS.filter(e => state.on.has(e));
  const labels = [];
  for (const e of on) {
    const ii = S.idx.filter(i => D.w.eqp[i] === e), r = regress(ii.map(i => D.w.p[i]), ii.map(yv));
    if (!isFinite(r.slope)) continue;
    const a = Math.max(xA, Math.min(...ii.map(i => D.w.p[i]))), b = Math.min(xB, Math.max(...ii.map(i => D.w.p[i])));
    el("line", { x1: x(a), y1: y(r.icpt + r.slope * a), x2: x(b), y2: y(r.icpt + r.slope * b), style: `stroke:${colorOf(e)};stroke-width:2` }, g);
    labels.push({ e, y: Math.min(Math.max(y(r.icpt + r.slope * b), m.t + 6), H - m.b), s: r.slope });
  }
  if (on.length <= 4) {
    labels.sort((a, b) => a.y - b.y); for (let k = 1; k < labels.length; k++) if (labels[k].y - labels[k - 1].y < 13) labels[k].y = labels[k - 1].y + 13;
    for (const L of labels) el("text", { class: "val", x: W - m.r + 6, y: L.y + 4, text: `${L.e} ${f(L.s, 2)}` }, svg);
  }
  const ring = el("circle", { r: 6, style: "fill:none;stroke:var(--ink);stroke-width:1.5;display:none" }, svg);
  const hit = el("rect", { x: m.l, y: m.t, width: W - m.l - m.r, height: H - m.t - m.b, style: "fill:transparent" }, svg);
  hit.addEventListener("pointermove", ev => {
    const rc = svg.getBoundingClientRect(), px = (ev.clientX - rc.left) * W / rc.width, py = (ev.clientY - rc.top) * H / rc.height;
    let best = -1, bd = 400;
    for (const i of S.idx) { const dx = x(D.w.p[i]) - px, dy = y(yv(i)) - py, dd = dx * dx + dy * dy; if (dd < bd) { bd = dd; best = i; } }
    if (best < 0) { ring.style.display = "none"; hideTip(); return; }
    ring.style.display = ""; ring.setAttribute("cx", x(D.w.p[best])); ring.setAttribute("cy", y(yv(best)));
    showTip(ev, `${D.w.eqp[best]} · ${D.w.wafer[best]}`, [
      [colorOf(D.w.eqp[best]), f(yv(best), 1) + " Å", state.y === "dev" ? "편차" : "두께"],
      [null, f(D.w.p[best], 2) + " hPa", "대기압"], [null, fmtDT(D.w.t[best]), "TKIN"]]);
  });
  hit.addEventListener("pointerleave", () => { ring.style.display = "none"; hideTip(); });
}

// ---------- 3. 설비별 민감도 ----------
function eqpStats(S) {
  return EQPS.filter(e => state.on.has(e)).map(e => {
    const ii = S.idx.filter(i => D.w.eqp[i] === e);
    return { e, ...regress(ii.map(i => D.w.p[i]), ii.map(i => D.w.thk[i])) };
  });
}
function renderSens(S) {
  const box = document.getElementById("sens"); box.replaceChildren();
  const st = eqpStats(S).filter(s => isFinite(s.slope));
  const W = box.clientWidth || 500, rowH = 30, m = { l: 80, r: 90, t: 6, b: 24 }, H = Math.max(120, m.t + m.b + st.length * rowH);
  const svg = el("svg", { viewBox: `0 0 ${W} ${H}`, height: H }, box);
  if (!st.length) { el("text", { x: W / 2, y: H / 2, "text-anchor": "middle", text: "계산할 수 있는 설비가 없습니다 (설비당 3장 이상 필요)" }, svg); return; }
  const lo = Math.min(0, ...st.map(s => s.slope)), hi = Math.max(0, ...st.map(s => s.slope)), p = (hi - lo) * 0.05 || 1;
  const [xA, xB] = rng("sens", "x", lo - (lo < 0 ? p : 0), hi + p), x = lin(xA, xB, m.l, W - m.r);
  const cl = v => Math.min(Math.max(v, m.l), W - m.r);
  for (const v of niceTicks(xA, xB, 5)) { el("line", { class: "grid", x1: x(v), x2: x(v), y1: m.t, y2: H - m.b }, svg); el("text", { x: x(v), y: H - 6, "text-anchor": "middle", text: tk(v) }, svg); }
  if (xA <= 0 && xB >= 0) el("line", { class: "base", x1: x(0), x2: x(0), y1: m.t, y2: H - m.b }, svg);
  const g = clipG(svg, m.l, m.t, W - m.l - m.r, H - m.t - m.b);
  st.forEach((s, k) => {
    const yc = m.t + k * rowH + rowH / 2, a = x(Math.min(0, s.slope)), b = x(Math.max(0, s.slope));
    el("text", { class: "lbl", x: m.l - 8, y: yc + 4, "text-anchor": "end", text: s.e }, svg);
    const bar = el("rect", { x: a, y: yc - 9, width: Math.max(2, b - a), height: 18, rx: 4, style: `fill:${colorOf(s.e)}` }, g);
    const sd = Math.abs(s.slope) < 0.1 ? 3 : 2, lx = s.slope >= 0 ? cl(b) + 6 : cl(x(0)) + 6;
    el("text", { class: "val", x: lx, y: yc + 4, "text-anchor": "start", text: `${f(s.slope, sd)}  (R² ${f(s.r2, 2)})` }, svg);
    const h = el("rect", { x: m.l, y: yc - rowH / 2, width: W - m.l - m.r, height: rowH, style: "fill:transparent" }, svg);
    h.addEventListener("pointermove", ev => { bar.style.opacity = 0.8; showTip(ev, s.e, [[colorOf(s.e), f(s.slope, 3) + " Å/hPa", "기울기"], [null, f(s.r2, 3), "R²"], [null, String(s.n), "웨이퍼 수"]]); });
    h.addEventListener("pointerleave", () => { bar.style.opacity = 1; hideTip(); });
  });
}

// ---------- 4. 대기압 구간별 ----------
function renderBins(S) {
  const box = document.getElementById("bins"); box.replaceChildren();
  const W = box.clientWidth || 500, H = 280, m = { l: 52, r: 12, t: 18, b: 40 }, bw = D.meta.bin;
  const svg = el("svg", { viewBox: `0 0 ${W} ${H}`, height: H }, box);
  const gm = new Map();
  for (const i of S.idx) { const k = Math.floor(D.w.p[i] / bw) * bw; if (!gm.has(k)) gm.set(k, []); gm.get(k).push(S.dev.get(i)); }
  const bins = [...gm.entries()].sort((a, b) => a[0] - b[0]).map(([k, v]) => {
    const n = v.length, mean = v.reduce((a, b) => a + b, 0) / n, sd = n > 1 ? Math.sqrt(v.reduce((a, b) => a + (b - mean) ** 2, 0) / (n - 1)) : 0;
    return { k, n, mean, sd };
  });
  if (!bins.length) { el("text", { x: W / 2, y: H / 2, "text-anchor": "middle", text: "선택된 데이터가 없습니다" }, svg); return; }
  const lo = Math.min(0, ...bins.map(b => b.mean - b.sd)), hi = Math.max(0, ...bins.map(b => b.mean + b.sd)), p = (hi - lo) * 0.08 || 1;
  const [yA, yB] = rng("bins", "y", lo - p, hi + p);
  const y = lin(yA, yB, H - m.b, m.t), slot = (W - m.l - m.r) / bins.length, bwid = Math.max(4, Math.min(48, slot - 6));
  const y0 = y(Math.min(Math.max(0, yA), yB));
  for (const v of niceTicks(yA, yB, 5)) { el("line", { class: "grid", x1: m.l, x2: W - m.r, y1: y(v), y2: y(v) }, svg); el("text", { x: m.l - 6, y: y(v) + 4, "text-anchor": "end", text: tk(v) }, svg); }
  el("line", { class: "base", x1: m.l, x2: W - m.r, y1: y0, y2: y0 }, svg);
  el("text", { class: "lbl", x: 4, y: m.t - 6, text: "Δ 두께 (Å)" }, svg);
  el("text", { class: "lbl", x: (m.l + W - m.r) / 2, y: H - 4, "text-anchor": "middle", text: `대기압 구간 시작값 (hPa, 폭 ${bw})` }, svg);
  const g = clipG(svg, m.l, m.t, W - m.l - m.r, H - m.t - m.b);
  const every = Math.ceil(bins.length / Math.max(1, Math.floor((W - m.l - m.r) / 44)));
  bins.forEach((b, k) => {
    const cx = m.l + slot * (k + 0.5), a = y(Math.max(0, b.mean)), z = y(Math.min(0, b.mean));
    const bar = el("rect", { x: cx - bwid / 2, y: a, width: bwid, height: Math.max(1, z - a), rx: 3, style: "fill:var(--s1)" }, g);
    el("line", { x1: cx, x2: cx, y1: y(b.mean - b.sd), y2: y(b.mean + b.sd), style: "stroke:var(--ink2);stroke-width:1.5" }, g);
    if (k % every === 0) el("text", { x: cx, y: H - m.b + 16, "text-anchor": "middle", text: f(b.k, bw < 1 ? 1 : 0) }, svg);
    const h = el("rect", { x: cx - slot / 2, y: m.t, width: slot, height: H - m.t - m.b, style: "fill:transparent" }, svg);
    h.addEventListener("pointermove", ev => { bar.style.opacity = 0.75; showTip(ev, `${f(b.k, 1)} ~ ${f(b.k + bw, 1)} hPa`, [["var(--s1)", f(b.mean, 2) + " Å", "평균 편차"], [null, f(b.sd, 2) + " Å", "σ"], [null, String(b.n), "웨이퍼 수"]]); });
    h.addEventListener("pointerleave", () => { bar.style.opacity = 1; hideTip(); });
  });
}

// ---------- 5. 시차 분석 ----------
function renderLag() {
  const box = document.getElementById("lag"); box.replaceChildren();
  const L = D.lag.filter(r => r.r != null);
  const W = box.clientWidth || 500, H = 280, m = { l: 52, r: 16, t: 18, b: 40 };
  const svg = el("svg", { viewBox: `0 0 ${W} ${H}`, height: H }, box);
  if (!L.length) { el("text", { x: W / 2, y: H / 2, "text-anchor": "middle", text: "데이터 없음" }, svg); return; }
  const xlo = L[0].lag, xhi = L[L.length - 1].lag, rs = L.map(r => r.r);
  const lo = Math.min(...rs), hi = Math.max(...rs), p = (hi - lo) * 0.15 || 0.05;
  const [yA, yB] = rng("lag", "y", lo - p, hi + p);
  const x = lin(xlo, xhi, m.l, W - m.r), y = lin(yA, yB, H - m.b, m.t);
  for (const v of niceTicks(yA, yB, 5)) { el("line", { class: "grid", x1: m.l, x2: W - m.r, y1: y(v), y2: y(v) }, svg); el("text", { x: m.l - 6, y: y(v) + 4, "text-anchor": "end", text: tk(v) }, svg); }
  for (const v of niceTicks(xlo, xhi, 6)) el("text", { x: x(v), y: H - m.b + 16, "text-anchor": "middle", text: tk(v) }, svg);
  el("line", { class: "base", x1: m.l, x2: W - m.r, y1: H - m.b, y2: H - m.b }, svg);
  el("text", { class: "lbl", x: (m.l + W - m.r) / 2, y: H - 4, "text-anchor": "middle", text: "시차 (분, TKIN_TIME 이전)" }, svg);
  el("text", { class: "lbl", x: 4, y: m.t - 6, text: "Pearson r" }, svg);
  const g = clipG(svg, m.l - 6, m.t, W - m.l - m.r + 12, H - m.t - m.b);
  el("path", { d: L.map((r, k) => (k ? "L" : "M") + x(r.lag) + "," + y(r.r)).join(""), style: "fill:none;stroke:var(--s1);stroke-width:2" }, g);
  const best = L.reduce((a, b) => Math.abs(b.r) > Math.abs(a.r) ? b : a);
  if (best.r >= yA && best.r <= yB) {
    el("circle", { cx: x(best.lag), cy: y(best.r), r: 5, style: "fill:var(--s1);stroke:var(--surface);stroke-width:2" }, svg);
    const right = x(best.lag) > W - 160;
    el("text", { class: "val", x: x(best.lag) + (right ? -8 : 8), y: y(best.r) - 8, "text-anchor": right ? "end" : "start", text: `최대 |r| ${f(best.r, 3)} @ ${best.lag}분` }, svg);
  }
  const cross = el("line", { y1: m.t, y2: H - m.b, style: "stroke:var(--muted);stroke-width:1;display:none" }, svg);
  const dot = el("circle", { r: 4, style: "fill:var(--s1);stroke:var(--surface);stroke-width:2;display:none" }, svg);
  const hit = el("rect", { x: m.l, y: m.t, width: W - m.l - m.r, height: H - m.t - m.b, style: "fill:transparent" }, svg);
  hit.addEventListener("pointermove", ev => {
    const rc = svg.getBoundingClientRect(), px = (ev.clientX - rc.left) * W / rc.width;
    const r = L.reduce((a, b) => Math.abs(x(b.lag) - px) < Math.abs(x(a.lag) - px) ? b : a);
    cross.style.display = dot.style.display = ""; cross.setAttribute("x1", x(r.lag)); cross.setAttribute("x2", x(r.lag));
    dot.setAttribute("cx", x(r.lag)); dot.setAttribute("cy", Math.min(Math.max(y(r.r), m.t), H - m.b));
    showTip(ev, `시차 ${r.lag}분`, [["var(--s1)", f(r.r, 3), "Pearson r"], [null, f(r.slope, 2) + " Å/hPa", "기울기"], [null, String(r.n), "웨이퍼 수"]]);
  });
  hit.addEventListener("pointerleave", () => { cross.style.display = dot.style.display = "none"; hideTip(); });
}

// ---------- 표 ----------
function renderTable(S) {
  const rows = [];
  const all = regress(S.idx.map(i => D.w.p[i]), S.idx.map(i => S.dev.get(i)));
  rows.push(["전체 (설비 평균 제거)", all]);
  for (const s of eqpStats(S)) rows.push(["설비 " + s.e, s]);
  const t = document.createElement("table"), hd = t.createTHead().insertRow();
  for (const h of ["구분", "N", "기울기 (Å/hPa)", "R²", "Pearson r", "두께 평균 (Å)", "두께 σ (Å)"]) { const th = document.createElement("th"); th.textContent = h; hd.appendChild(th); }
  const tb = t.createTBody();
  for (const [name, r] of rows) {
    const tr = tb.insertRow();
    for (const v of [name, String(r.n), f(r.slope, 3), f(r.r2, 3), f(r.r, 3), name.startsWith("전체") ? "-" : f(r.mean, 1), name.startsWith("전체") ? "-" : f(r.sd, 2)]) tr.insertCell().textContent = v;
  }
  document.getElementById("tbl").replaceChildren(t);
}

// ---------- 전체 그리기 ----------
function render() {
  const S = selected();
  const rl = document.getElementById("rangeLabel"), rb = document.getElementById("rangeReset");
  rl.textContent = state.range ? `기간: ${fmtDT(state.range[0])} ~ ${fmtDT(state.range[1])}` : "";
  rb.hidden = !state.range;
  renderKpis(S); renderTS(S); renderScatter(S); renderSens(S); renderBins(S); renderLag(); renderTable(S);
  updatePlaceholders();
}
let rt; addEventListener("resize", () => { clearTimeout(rt); rt = setTimeout(render, 150); });
render();
</script>
</body>
</html>
'''


# ============================================================
# 엑셀
# ============================================================
def write_excel(path, df, unmatched, summary, bins, lag, charts):
    cm, cp = COLUMNS["meas"], COLUMNS["proc"]
    names = {"lot": cm["lot"], "wafer": cm["wafer"], "eqp": cp["eqp"], "time": cp["time"],
             "meas_time": cm.get("meas_time") or "계측시간", "thickness": "두께(Å)"}
    data = df.rename(columns={
        **names, "pressure": "대기압(hPa)", "pressure_time": "매칭된 대기압 시간",
        f"pressure_avg_{WINDOW_MIN}m": f"직전{WINDOW_MIN}분 평균 대기압",
        f"pressure_change_{WINDOW_MIN}m": f"{WINDOW_MIN}분간 대기압 변화",
        "eqp_mean": "설비 평균 두께", "thickness_dev": "설비 평균 대비 편차(Å)"})
    um = unmatched.rename(columns=names)
    settings = pd.DataFrame({"항목": [
        "쿼리2-3 연결 키", "대기압 매칭 기준 시간", "ITEM_ID 필터", "매칭 허용 오차(분)",
        "평균/변화량 구간(분)", "대기압 구간 폭(hPa)", "매칭 웨이퍼 수", "매칭 실패 웨이퍼 수"], "값": [
        "LOT+WAFER" if cp.get("wafer") else "LOT", cp["time"], ITEM_FILTER,
        MATCH_TOLERANCE_MIN, WINDOW_MIN, PRESSURE_BIN_HPA, len(df), len(unmatched)]})
    sheets = {"요약_회귀": summary, "대기압구간별": bins, "시차분석": lag,
              "매칭데이터": data, "매칭실패": um, "설정": settings}
    if len(data) > EXCEL_MAX_ROWS:  # 행이 많으면 원본 행은 CSV 로 (엑셀로 바로 열림)
        csv_path = Path(path).with_name("matched_data.csv")
        data.to_csv(csv_path, index=False, encoding="utf-8-sig", date_format="%Y-%m-%d %H:%M:%S")
        sheets["매칭데이터"] = pd.DataFrame({"안내": [
            f"매칭 데이터가 {len(data):,}행으로 많아서 {csv_path.name} 파일로 따로 저장했습니다.",
            f"(EXCEL_MAX_ROWS = {EXCEL_MAX_ROWS:,} 설정으로 기준을 바꿀 수 있습니다)"]})
        print(f"  매칭데이터 {len(data):,}행 -> {csv_path}")
    # 숫자는 미리 반올림 (셀마다 서식을 지정하면 행이 많을 때 매우 느림)
    for name, d in sheets.items():
        d = d.copy()
        num = d.select_dtypes("float").columns
        d[num] = d[num].round(6 if name in ("요약_회귀", "시차분석") else 4)
        sheets[name] = d

    def width(d, col):
        sample = d[col].head(200).astype(str).str.len().max() if len(d) else 0
        return min(max(10, max(len(str(col)), sample or 0) * 1.3), 40)

    try:
        import xlsxwriter  # noqa: F401  (있으면 openpyxl 보다 훨씬 빠름. Anaconda 에 기본 포함)
        engine = "xlsxwriter"
    except ImportError:
        engine = "openpyxl"
    with pd.ExcelWriter(path, engine=engine, datetime_format="yyyy-mm-dd hh:mm:ss") as xw:
        for name, d in sheets.items():
            d.to_excel(xw, sheet_name=name, index=False)
        if engine == "xlsxwriter":
            for name, d in sheets.items():
                ws = xw.sheets[name]
                ws.freeze_panes(1, 0)
                for i, col in enumerate(d.columns):
                    ws.set_column(i, i, width(d, col))
            ws = xw.book.add_worksheet("그래프")
            row = 0
            for f in charts:
                ws.insert_image(row, 0, str(f))
                row += 24
            xw.book.worksheets_objs.insert(0, xw.book.worksheets_objs.pop())  # 그래프 시트를 맨 앞으로
            ws.activate()
        else:
            from openpyxl.drawing.image import Image
            from openpyxl.utils import get_column_letter
            for name, d in sheets.items():
                ws = xw.sheets[name]
                ws.freeze_panes = "A2"
                for i, col in enumerate(d.columns, 1):
                    ws.column_dimensions[get_column_letter(i)].width = width(d, col)
            ws = xw.book.create_sheet("그래프", 0)
            row = 1
            for f in charts:
                img = Image(str(f))
                ws.add_image(img, f"A{row}")
                row += int(img.height / 20) + 2


# ============================================================
# main
# ============================================================
class Clock:
    """단계별 소요 시간 기록 (어디가 느린지 확인용)."""
    def __init__(self):
        import time
        self._time = time.perf_counter
        self.t = self._time()
        self.steps = []

    def __call__(self, name):
        now = self._time()
        self.steps.append((name.strip(), now - self.t))
        print(f"  ({name.strip()}: {now - self.t:.1f}초)")
        self.t = now

    def summary(self):
        total = sum(s for _, s in self.steps)
        return f"총 {total:.1f}초 = " + ", ".join(f"{n} {s:.1f}" for n, s in self.steps)


def main(argv=None):
    ap = argparse.ArgumentParser(description="대기압 vs HSOH 두께 분석")
    ap.add_argument("--source", choices=["query", "csv"], default=DATA_SOURCE,
                    help=f"데이터 가져오는 방법 (기본: {DATA_SOURCE})")
    ap.add_argument("--from", dest="date_from", metavar="YYYY-MM-DD",
                    help=f"분석 시작일, 계측 기준 (조회 모드 기본: {DATE_FROM}, csv 모드 기본: 전체)")
    ap.add_argument("--to", dest="date_to", metavar="YYYY-MM-DD",
                    help=f"분석 종료일, 포함 (조회 모드 기본: {DATE_TO}, csv 모드 기본: 전체)")
    for kind, folder in (("pressure", "pressure"), ("meas", "meas"), ("proc", "proc")):
        ap.add_argument(f"--{kind}", default=ROOT / "input" / folder,
                        help=f"csv 모드: {SOURCES[kind]} CSV 파일 또는 폴더 (기본: input/{folder})")
    ap.add_argument("--out", default=ROOT / "output", help="결과 폴더 (기본: output)")
    args, _ = ap.parse_known_args(argv)  # 주피터 커널 인자는 무시

    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    clock = Clock()

    if args.source == "query":
        args.date_from = args.date_from or DATE_FROM
        args.date_to = args.date_to or DATE_TO
    start = end = None
    if args.date_from or args.date_to:
        try:
            start = pd.Timestamp(args.date_from).normalize() if args.date_from else None
            end = (pd.Timestamp(args.date_to).normalize() + pd.Timedelta(days=1)
                   if args.date_to else None)
        except ValueError:
            raise SystemExit("[오류] 날짜는 YYYY-MM-DD 형식으로 입력하세요.")
        if start is not None and end is not None and end <= start:
            raise SystemExit("[오류] 종료일이 시작일보다 빠릅니다.")

    if args.source == "query":
        if start is None or end is None:
            raise SystemExit("[오류] 조회 모드에서는 DATE_FROM / DATE_TO 가 필요합니다.")
        print(f"[1/5] 조회 ({start:%Y-%m-%d} ~ {end - pd.Timedelta(days=1):%Y-%m-%d})")
        raw_p, raw_m, raw_c = fetch_from_query(start, end, out_dir)
    else:
        print("[1/5] CSV 읽기")
        raw_p = read_all(args.pressure, "pressure")
        raw_m = read_all(args.meas, "meas")
        raw_c = read_all(args.proc, "proc")
    pressure = load_pressure(raw_p)
    meas = load_meas(raw_m)
    proc = load_proc(raw_c)
    if "meas_time" in meas and (start is not None or end is not None):
        mt = meas["meas_time"]
        keep = pd.Series(True, index=meas.index)
        if start is not None:
            keep &= mt.isna() | (mt >= start)
        if end is not None:
            keep &= mt.isna() | (mt < end)
        meas = meas[keep].reset_index(drop=True)
    if pressure.empty or meas.empty or proc.empty:
        raise SystemExit("[오류] 사용할 수 있는 데이터가 없습니다. "
                         f"(대기압 {len(pressure)} / 계측 {len(meas)} / 공정 {len(proc)}행)")

    clock("1/5 조회/읽기")
    print("[2/5] 계측 + 공정 연결 (LOT" + (", WAFER" if "wafer" in proc else "") + ")")
    soh, no_proc = join_meas_proc(meas, proc)
    print(f"  연결 {len(soh)} / 공정정보 없음 {len(no_proc)} (계측 웨이퍼 {len(meas)})")
    if soh.empty:
        raise SystemExit("[오류] 계측과 공정 정보가 연결된 웨이퍼가 없습니다. "
                         "두 쿼리의 LOT 값 형식이 같은지 확인하세요.")
    print(f"  대기압 {pressure['time'].min()} ~ {pressure['time'].max()} ({len(pressure)}개)")
    print(f"  TKIN   {soh['time'].min()} ~ {soh['time'].max()} ({len(soh)}웨이퍼)")

    clock("2/5 연결")
    print("[3/5] 대기압 매칭 (TKIN_TIME 기준)")
    merged = match(soh, pressure)
    cols = ["lot", "wafer", "eqp", "time", "meas_time", "thickness"]
    no_p = merged[merged["pressure"].isna()].assign(사유="대기압 없음")
    unmatched = pd.concat([no_proc.assign(사유="공정정보 없음"), no_p], ignore_index=True)
    unmatched = unmatched[[c for c in cols if c in unmatched] + ["사유"]]
    df = merged[merged["pressure"].notna()].copy()
    df["eqp_mean"] = df.groupby("eqp")["thickness"].transform("mean")
    df["thickness_dev"] = df["thickness"] - df["eqp_mean"]
    print(f"  매칭 {len(df)} / 대기압 없음 {len(no_p)} (허용 ±{MATCH_TOLERANCE_MIN}분)")
    if df.empty:
        raise SystemExit("[오류] 대기압이 매칭된 웨이퍼가 없습니다. 두 데이터의 기간을 확인하세요.")

    clock("3/5 매칭")
    print("[4/5] 분석")
    summary = summary_table(df)
    bins = bin_table(df)
    lag = lag_table(df, pressure)

    clock("4/5 분석")
    print("[5/5] 결과 저장")
    charts = make_charts(df, pressure, summary, lag, out_dir) if MAKE_PNG else []
    clock("  그래프(PNG)")
    html = out_dir / "pressure_vs_soh_report.html"
    write_html(html, df, pressure, lag)
    clock("  HTML")
    xlsx = out_dir / "pressure_vs_soh_report.xlsx"
    try:
        write_excel(xlsx, df, unmatched, summary, bins, lag, charts)
    except PermissionError:
        raise SystemExit(f"[오류] {xlsx} 가 열려 있어 저장할 수 없습니다. 엑셀을 닫고 다시 실행하세요.")
    clock("  엑셀")

    pd.set_option("display.width", 200)
    print("\n=== 대기압 1 hPa 당 두께 변화 ===")
    print(summary[["구분", "N", "기울기(Å/hPa)", "R²", "p-value"]]
          .to_string(index=False, float_format=lambda v: f"{v:.4g}"))
    print(f"\n소요 시간: {clock.summary()}")
    print(f"결과: {html}  <- 브라우저로 열기 (인터랙티브)")
    print(f"      {xlsx}")
    for f in charts:
        print(f"      {f}")


if __name__ == "__main__":
    main()
